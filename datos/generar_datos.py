"""Genera los 3 extractos bancarios ficticios del ejercicio en la carpeta datos/.
Datos inventados: Entertainment Forever es una empresa ficticia.
"""
import csv, datetime as dt, os
os.chdir(os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment
D=dt.date
# ---------- Santander (ES): ; separator, decimal comma, dd/mm/yyyy
san=[(D(2026,10,20),'ES1200491500001234567890','Cobros Stripe liquidación',245320.50,'EUR'),
 (D(2026,10,20),'ES1200491500001234567890','Pago promotor Candlelight Madrid',-38400.00,'EUR'),
 (D(2026,10,21),'ES1200491500001234567890','Nóminas octubre',-182750.00,'EUR'),
 (D(2026,10,21),'ES1200491500001234567890','Transferencia intragrupo a Entertainment Forever UK Ltd',-50000.00,'EUR'),
 (D(2026,10,21),'ES4500491500009876543210','Cobro Adyen USD',54210.00,'USD'),
 (D(2026,10,22),'ES1200491500001234567890','Comisión transferencia internacional',-125.40,'EUR'),
 (D(2026,10,22),'ES4500491500009876543210','Pago proveedor tecnología',-12500.00,'USD'),
 (D(2026,10,23),'ES1200491500001234567890','Cobros Stripe liquidación',198045.75,'EUR'),
 (D(2026,10,23),'ES1200491500001234567890','Pago recinto IFEMA',-61200.00,'EUR'),
 (D(2026,10,24),'ES4500491500009876543210','Comisión mantenimiento cuenta',-45.00,'USD'),
 (D(2026,10,24),'ES1200491500001234567890','Liquidación IVA Hacienda',-74310.20,'EUR')]
def es(n):
    s=f"{abs(n):,.2f}".replace(',','X').replace('.',',').replace('X','.')
    return ('-' if n<0 else '')+s
with open('Santander_ES.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f,delimiter=';')
    w.writerow(['Fecha','Cuenta','Concepto','Importe','Divisa'])
    for d,a,c,i,m in san: w.writerow([d.strftime('%d/%m/%Y'),a,c,es(i),m])
# ---------- Barclays (UK): xlsx, title rows, Debit/Credit columns
bar=[(D(2026,10,20),'GB29BARC20000012345678','Stripe settlement',None,132450.00,'GBP'),
 (D(2026,10,20),'GB29BARC20000012345678','Payment to venue - Kew Gardens',22800.00,None,'GBP'),
 (D(2026,10,21),'GB82BARC20000087654321','Intercompany transfer from Entertainment Forever Espana SL',None,50000.00,'EUR'),
 (D(2026,10,21),'GB29BARC20000012345678','Payroll October',96300.00,None,'GBP'),
 (D(2026,10,22),'GB29BARC20000012345678','Bank charges',85.60,None,'GBP'),
 (D(2026,10,23),'GB82BARC20000087654321','Payment to promoter - Dublin',18750.00,None,'EUR'),
 (D(2026,10,23),'GB29BARC20000012345678','Stripe settlement',None,87925.40,'GBP'),
 (D(2026,10,24),'GB29BARC20000012345678','HMRC VAT payment',41200.00,None,'GBP')]
wb=Workbook(); ws=wb.active; ws.title='Statement'
ws['A1']='Barclays Business - Account statement'; ws['A1'].font=Font(bold=True,size=13)
ws['A2']='Period: 20/10/2026 - 24/10/2026'
hdr=['Date','Account','Description','Debit','Credit','Currency']
for j,h in enumerate(hdr,1): ws.cell(row=4,column=j,value=h).font=Font(bold=True)
for i,(d,a,c,db,cr,m) in enumerate(bar,5):
    for j,v in enumerate([d,a,c,db,cr,m],1):
        cell=ws.cell(row=i,column=j,value=v)
        if j==1: cell.number_format='DD/MM/YYYY'
        if j in (4,5): cell.number_format='#,##0.00'
for col,wid in zip('ABCDEF',[12,26,44,12,12,10]): ws.column_dimensions[col].width=wid
wb.save('Barclays_UK.xlsx')
# ---------- Chase (US): comma, MM/DD/YYYY, signed amounts, running balance, no account/currency
ch=[(D(2026,10,20),'STRIPE PAYOUT',176880.25),(D(2026,10,21),'VENUE PAYMENT CANDLELIGHT NYC',-45600.00),
 (D(2026,10,21),'PAYROLL ADP',-238400.00),(D(2026,10,22),'WIRE TRANSFER FEE',-35.00),
 (D(2026,10,23),'STRIPE PAYOUT',143215.60),(D(2026,10,24),'PROMOTER PAYMENT CHICAGO',-29950.00)]
bal=1250000.00
with open('Chase_US.csv','w',newline='',encoding='utf-8') as f:
    w=csv.writer(f)
    w.writerow(['Posting Date','Description','Amount','Type','Balance'])
    for d,c,a in ch:
        bal=round(bal+a,2)
        w.writerow([d.strftime('%m/%d/%Y'),c,f"{a:.2f}",'CREDIT' if a>0 else 'DEBIT',f"{bal:.2f}"])
print('Extractos generados: Santander_ES.csv, Barclays_UK.xlsx, Chase_US.csv')
